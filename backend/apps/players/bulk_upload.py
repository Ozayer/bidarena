from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.workbook import Workbook

from apps.tournaments.models import Position

from .models import Player

REQUIRED_COLUMNS = ['name', 'base price']
KNOWN_COLUMNS = ['name', 'position', 'base price']


def _normalize_header(value):
    return str(value or '').strip().lower().replace('_', ' ')


def build_template_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Players'
    ws.append(['Name', 'Position', 'Base Price', 'Age'])
    ws.append(['Jane Doe', 'Defender', 100, 27])
    return wb


def parse_and_create_players(file, tournament):
    """Parses an uploaded Excel file and creates valid players, row by row.

    Returns {'created': int, 'total_rows': int, 'errors': [{'row': int, 'errors': [str]}]}.
    Invalid rows are skipped (not created) but don't block valid rows.
    """
    wb = openpyxl.load_workbook(file, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {'created': 0, 'total_rows': 0, 'errors': [{'row': 1, 'errors': ['File is empty.']}]}

    header = [_normalize_header(c) for c in rows[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        return {
            'created': 0,
            'total_rows': 0,
            'errors': [{'row': 1, 'errors': [f"Missing required column: {c.title()}" for c in missing]}],
        }

    col_index = {name: idx for idx, name in enumerate(header) if name}
    extra_columns = [c for c in header if c and c not in KNOWN_COLUMNS]

    positions_by_name = {
        p.name.strip().lower(): p for p in Position.objects.filter(tournament=tournament)
    }

    created = 0
    errors = []
    data_rows = rows[1:]

    for offset, row in enumerate(data_rows):
        row_number = offset + 2  # account for header row + 1-indexing
        if row is None or all(cell in (None, '') for cell in row):
            continue

        row_errors = []

        name = row[col_index['name']] if col_index['name'] < len(row) else None
        name = str(name).strip() if name is not None else ''
        if not name:
            row_errors.append('Name is required.')

        base_price_raw = row[col_index['base price']] if col_index['base price'] < len(row) else None
        base_price = None
        if base_price_raw is None or str(base_price_raw).strip() == '':
            row_errors.append('Base price is required.')
        else:
            try:
                base_price = Decimal(str(base_price_raw))
                if base_price < 0:
                    row_errors.append('Base price must be a non-negative number.')
            except InvalidOperation:
                row_errors.append(f"Base price '{base_price_raw}' is not a valid number.")

        position = None
        if 'position' in col_index and col_index['position'] < len(row):
            position_raw = row[col_index['position']]
            position_name = str(position_raw).strip() if position_raw is not None else ''
            if position_name:
                position = positions_by_name.get(position_name.lower())
                if position is None:
                    row_errors.append(f"Unknown position '{position_name}'.")

        if row_errors:
            errors.append({'row': row_number, 'errors': row_errors})
            continue

        extra_info = {}
        for col in extra_columns:
            idx = col_index[col]
            if idx < len(row) and row[idx] not in (None, ''):
                extra_info[col] = row[idx]

        Player.objects.create(
            tournament=tournament,
            name=name,
            position=position,
            base_price=base_price,
            extra_info=extra_info,
        )
        created += 1

    return {'created': created, 'total_rows': len(data_rows), 'errors': errors}
